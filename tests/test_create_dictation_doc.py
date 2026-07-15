import importlib.util
import tempfile
import unittest
import zipfile
from pathlib import Path

from docx import Document


ROOT = Path(__file__).resolve().parents[1]
SCRIPT = ROOT / "scripts" / "create_dictation_doc.py"


def load_module():
    spec = importlib.util.spec_from_file_location("create_dictation_doc", SCRIPT)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


class EnglishDictationGeneratorTests(unittest.TestCase):
    def test_reads_five_column_a4_contract_from_bundled_html(self):
        module = load_module()
        contract = module.load_template_contract()
        self.assertEqual(contract["columns"], 5)
        self.assertEqual(contract["accent_color"], "C0392B")
        self.assertEqual(contract["title"], "看中文写英文")

    def test_rejects_empty_items(self):
        module = load_module()
        with self.assertRaisesRegex(ValueError, "至少提供一个英语默写词条"):
            module.validate_items([])

    def test_rejects_prompt_with_english_metadata(self):
        module = load_module()
        with self.assertRaisesRegex(ValueError, "中文提示不得包含英文字母"):
            module.validate_items([{"prompt": "n. 星期一", "answer": "Monday"}])

    def test_preserves_source_order(self):
        module = load_module()
        items = [
            {"prompt": "星期一", "answer": "Monday"},
            {"prompt": "星期二", "answer": "Tuesday"},
            {"prompt": "星期三", "answer": "Wednesday"},
        ]
        self.assertEqual(module.validate_items(items), items)

    def test_generates_a4_five_column_student_doc_without_answers(self):
        module = load_module()
        items = [
            {"prompt": "星期一", "answer": "Monday"},
            {"prompt": "星期二", "answer": "Tuesday"},
            {"prompt": "通常", "answer": "usually"},
            {"prompt": "有时", "answer": "sometimes"},
            {"prompt": "总是", "answer": "always"},
            {"prompt": "经常", "answer": "often"},
        ]

        with tempfile.TemporaryDirectory() as tmp:
            output = Path(tmp) / "dictation.docx"
            module.create_document(items, output)
            doc = Document(output)
            section = doc.sections[0]
            self.assertAlmostEqual(section.page_width.cm, 21.0, places=1)
            self.assertAlmostEqual(section.page_height.cm, 29.7, places=1)
            self.assertEqual(len(doc.tables[-1].columns), 5)
            self.assertEqual(len(doc.tables[-1].rows), 2)
            self.assertEqual(len(doc.inline_shapes), len(items) + 2)
            with zipfile.ZipFile(output) as package:
                xml = package.read("word/document.xml").decode("utf-8")

        for prompt in ("星期一", "星期二", "通常", "有时", "总是", "经常"):
            self.assertIn(prompt, xml)
        for answer in ("Monday", "Tuesday", "usually", "sometimes", "always", "often"):
            self.assertNotIn(answer, xml)
        self.assertEqual(xml.count("<w:t>________</w:t>"), 6)
        self.assertIn("看中文写英文", xml)
        self.assertIn("共6题", xml)


if __name__ == "__main__":
    unittest.main()
