#!/usr/bin/env python3
"""Create an A4 English dictation DOCX from Chinese prompts."""

from __future__ import annotations

import argparse
import json
import math
import re
import tempfile
from pathlib import Path
from typing import Sequence

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont


TEMPLATE_PATH = Path(__file__).resolve().parents[1] / "assets" / "template.html"
MAX_ROWS_PER_PAGE = 13
ANSWER_LINE = "________"
FONT_CN = "Arial Unicode MS"
FONT_CANDIDATES = (
    "/Library/Fonts/Arial Unicode.ttf",
    "/System/Library/Fonts/STHeiti Medium.ttc",
    "/System/Library/Fonts/Supplemental/Songti.ttc",
    "/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc",
)


def load_template_contract() -> dict:
    html = TEMPLATE_PATH.read_text(encoding="utf-8")
    patterns = {
        "accent_color": r"--accent:\s*#([0-9a-fA-F]{6})",
        "columns": r"\.word-grid\{[^}]*grid-template-columns:\s*repeat\((\d+),",
        "title": r'<div class="sec-title">\s*([^<{]+)',
    }
    contract = {}
    for key, pattern in patterns.items():
        match = re.search(pattern, html)
        if not match:
            raise ValueError(f"模板缺少必要样式：{key}")
        contract[key] = match.group(1).strip()
    contract["columns"] = int(contract["columns"])
    contract["accent_color"] = contract["accent_color"].upper()
    if "@page{size:A4portrait" not in re.sub(r"\s+", "", html):
        raise ValueError("模板必须声明 A4 纵向打印")
    if "{{COUNT}}" not in html or "{{WORDS}}" not in html:
        raise ValueError("模板必须包含 COUNT 和 WORDS 占位符")
    return contract


TEMPLATE_CONTRACT = load_template_contract()


def validate_items(items: Sequence[dict]) -> list[dict]:
    if not items:
        raise ValueError("至少提供一个英语默写词条")
    validated = []
    for index, item in enumerate(items, start=1):
        if not isinstance(item, dict):
            raise ValueError(f"第{index}项必须是对象")
        prompt = item.get("prompt")
        answer = item.get("answer")
        if not isinstance(prompt, str) or not prompt.strip():
            raise ValueError(f"第{index}项缺少中文提示")
        if re.search(r"[A-Za-z]", prompt):
            raise ValueError(f"第{index}项中文提示不得包含英文字母")
        if not isinstance(answer, str) or not answer.strip():
            raise ValueError(f"第{index}项缺少英语答案")
        validated.append(item)
    return validated


def _set_cell_margins(cell, top=80, start=70, bottom=80, end=70):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def _set_table_borders(table, **edges):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge, attrs in edges.items():
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        for key, value in attrs.items():
            node.set(qn(f"w:{key}"), str(value))


def _set_table_geometry(table, widths: Sequence[int]):
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(width))
        grid.append(grid_col)
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            tc_w = cell._tc.get_or_add_tcPr().find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                cell._tc.get_or_add_tcPr().append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")


def _set_run_font(run, size: float, *, bold=False, color="1A1A1A", name=FONT_CN):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.rFonts
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.insert(0, r_fonts)
    for key in ("ascii", "hAnsi", "eastAsia"):
        r_fonts.set(qn(f"w:{key}"), name)


def _find_cjk_font() -> Path:
    for candidate in FONT_CANDIDATES:
        path = Path(candidate)
        if path.exists():
            return path
    raise RuntimeError("未找到可用于中文提示的字体")


def _set_picture_alt(shape, description: str):
    shape._inline.docPr.set("descr", description)


def _build_meta_image(directory: Path) -> Path:
    font = ImageFont.truetype(str(_find_cjk_font()), 27)
    canvas = Image.new("RGB", (1600, 130), "white")
    draw = ImageDraw.Draw(canvas)
    draw.line((0, 5, 1600, 5), fill="#1a1a1a", width=5)
    draw.line((0, 124, 1600, 124), fill="#1a1a1a", width=5)
    labels = ["姓名：____________", "班级：__________", "日期：______月______日", "得分：__________"]
    centers = [205, 570, 1020, 1420]
    for label, center in zip(labels, centers):
        box = draw.textbbox((0, 0), label, font=font)
        draw.text((center - (box[2] - box[0]) / 2, 46), label, font=font, fill="#1a1a1a")
    path = directory / "meta.png"
    canvas.save(path, dpi=(300, 300))
    return path


def _build_section_image(directory: Path, count: int) -> Path:
    title_font = ImageFont.truetype(str(_find_cjk_font()), 30)
    count_font = ImageFont.truetype(str(_find_cjk_font()), 23)
    canvas = Image.new("RGB", (1600, 92), "white")
    draw = ImageDraw.Draw(canvas)
    draw.rectangle((0, 15, 13, 77), fill=f"#{TEMPLATE_CONTRACT['accent_color']}")
    draw.text((36, 24), TEMPLATE_CONTRACT["title"], font=title_font, fill="#1a1a1a")
    draw.text((300, 31), f"共{count}题", font=count_font, fill="#888888")
    path = directory / "section.png"
    canvas.save(path, dpi=(300, 300))
    return path


def _build_prompt_image(directory: Path, prompt: str, index: int) -> Path:
    font = ImageFont.truetype(str(_find_cjk_font()), 38)
    canvas = Image.new("RGB", (420, 86), "white")
    draw = ImageDraw.Draw(canvas)
    box = draw.textbbox((0, 0), prompt, font=font)
    width = box[2] - box[0]
    height = box[3] - box[1]
    draw.text(((420 - width) / 2, (86 - height) / 2 - box[1]), prompt, font=font, fill="#1a1a1a")
    path = directory / f"prompt-{index:04d}.png"
    canvas.save(path, dpi=(300, 300))
    return path


def _add_meta_bar(doc, directory: Path):
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(4)
    shape = paragraph.add_run().add_picture(str(_build_meta_image(directory)), width=Cm(18.6))
    _set_picture_alt(shape, "姓名、班级、日期、得分填写栏")


def _add_section_title(doc, count: int, directory: Path):
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(4)
    shape = paragraph.add_run().add_picture(str(_build_section_image(directory, count)), width=Cm(18.6))
    _set_picture_alt(shape, f"{TEMPLATE_CONTRACT['title']} 共{count}题")


def _add_content_table(doc, items: Sequence[dict], directory: Path, start_index: int):
    columns = TEMPLATE_CONTRACT["columns"]
    rows = math.ceil(len(items) / columns)
    table = doc.add_table(rows=rows, cols=columns)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    _set_table_geometry(table, [2109] * columns)
    none = {"val": "nil"}
    _set_table_borders(table, top=none, bottom=none, left=none, right=none, insideH=none, insideV=none)
    for row in table.rows:
        cant_split = OxmlElement("w:cantSplit")
        row._tr.get_or_add_trPr().append(cant_split)
    for index, item in enumerate(items):
        cell = table.cell(index // columns, index % columns)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        _set_cell_margins(cell, top=140, start=70, bottom=140, end=70)
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(2)
        prompt_path = _build_prompt_image(directory, item["prompt"], start_index + index)
        shape = paragraph.add_run().add_picture(str(prompt_path), width=Cm(3.2))
        _set_picture_alt(shape, f"中文提示：{item['prompt']}")
        line = cell.add_paragraph()
        line.alignment = WD_ALIGN_PARAGRAPH.CENTER
        line.paragraph_format.space_before = Pt(0)
        line.paragraph_format.space_after = Pt(0)
        _set_run_font(line.add_run(ANSWER_LINE), 13, name="Arial")
    return table


def _add_page(doc, items: Sequence[dict], total_count: int, directory: Path, start_index: int):
    _add_meta_bar(doc, directory)
    _add_section_title(doc, total_count, directory)
    _add_content_table(doc, items, directory, start_index)


def create_document(items: Sequence[dict], output_path: Path | str):
    validated = validate_items(items)
    columns = TEMPLATE_CONTRACT["columns"]
    items_per_page = columns * MAX_ROWS_PER_PAGE
    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.2)
    section.bottom_margin = Cm(1.2)
    section.left_margin = Cm(1.2)
    section.right_margin = Cm(1.2)
    normal = doc.styles["Normal"]
    normal.font.name = FONT_CN
    normal.font.size = Pt(10.5)

    output = Path(output_path)
    output.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory() as tmp:
        image_dir = Path(tmp)
        for page_index, start in enumerate(range(0, len(validated), items_per_page)):
            if page_index:
                doc.add_section(WD_SECTION.NEW_PAGE)
            _add_page(doc, validated[start:start + items_per_page], len(validated), image_dir, start)
        doc.save(output)
    return output


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("input_json", type=Path)
    parser.add_argument("output_docx", type=Path)
    args = parser.parse_args()
    payload = json.loads(args.input_json.read_text(encoding="utf-8"))
    create_document(payload.get("items", []), args.output_docx)


if __name__ == "__main__":
    main()
